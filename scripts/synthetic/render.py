"""Renderers for the synthetic corpus: one block model -> PDF / DOCX / XLSX / MD.

Every output carries the FICTION banner on every page (PDF header, DOCX header, XLSX banner row +
README sheet, MD front matter). Content is authored in the content_*.py modules as documents made of
blocks:  ("h", text) heading · ("p", text) paragraph · ("b", [items]) bullets ·
         ("t", headers, rows) table · ("note", text) call-out
XLSX documents instead carry `sheets`: [(name, headers, rows, col_widths?)].
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

BANNER = ("FICTIONAL DOCUMENT — FOR EDUCATIONAL PURPOSES ONLY — NOT REAL DATA. "
          "Not an actual Kohler Co. policy or record. Created for an academic AI prototype; "
          "all figures, names, thresholds and identifiers are invented.")
BANNER_SHORT = "FICTIONAL — EDUCATIONAL PURPOSES ONLY — NOT REAL DATA"
RED = colors.HexColor("#B00020")

# Helvetica (reportlab default) has no rupee glyph; Arial does. Register once, with a bold face for <b> tags.
_FONT_DIR = Path("C:/Windows/Fonts")
if (_FONT_DIR / "arial.ttf").exists():
    pdfmetrics.registerFont(TTFont("Arial", str(_FONT_DIR / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("Arial-Bold", str(_FONT_DIR / "arialbd.ttf")))
    from reportlab.lib.fonts import addMapping
    addMapping("Arial", 0, 0, "Arial")
    addMapping("Arial", 1, 0, "Arial-Bold")
    addMapping("Arial", 0, 1, "Arial")
    addMapping("Arial", 1, 1, "Arial-Bold")
    BASE_FONT, BOLD_FONT = "Arial", "Arial-Bold"
else:  # fallback keeps the pipeline working on machines without Arial
    BASE_FONT, BOLD_FONT = "Helvetica", "Helvetica-Bold"


def _wrap(text: str, width: int) -> list[str]:
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width and cur:
            lines.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        lines.append(cur)
    return lines


def _meta_line(doc: dict) -> str:
    bits = [f"Doc code: {doc.get('doc_code', doc['id'])}"]
    if doc.get("version"):
        bits.append(f"Version {doc['version']}")
    if doc.get("effective_date"):
        bits.append(f"Effective {doc['effective_date']}")
    if doc.get("supersedes"):
        bits.append(f"Supersedes {doc['supersedes']}")
    bits.append(f"Classification: {doc['sensitivity'].upper()}")
    bits.append(f"Audience: {doc['audience']}")
    bits.append(f"Region: {doc['region']}")
    if doc.get("owner"):
        bits.append(f"Owner: {doc['owner']}")
    return " · ".join(bits)


# ───────────────────────────── PDF ─────────────────────────────
def render_pdf(doc: dict, out: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=BASE_FONT, fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontName=BOLD_FONT, fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName=BOLD_FONT, fontSize=11.5, leading=14, spaceBefore=9, spaceAfter=3)
    meta = ParagraphStyle("meta", parent=body, fontSize=8, textColor=colors.HexColor("#444444"), spaceAfter=8)
    note = ParagraphStyle("note", parent=body, backColor=colors.HexColor("#FFF4E5"), borderPadding=4, leftIndent=4)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11, spaceAfter=0)

    def header_footer(canvas, d):
        canvas.saveState()
        w, h = A4
        canvas.setFillColor(RED)
        canvas.rect(15 * mm, h - 17 * mm, w - 30 * mm, 9 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(BOLD_FONT, 7.2)
        b1, b2 = (_wrap(BANNER, 118) + ["", ""])[:2]
        canvas.drawCentredString(w / 2, h - 11.2 * mm, b1)
        canvas.drawCentredString(w / 2, h - 14.6 * mm, b2)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.setFont(BASE_FONT, 7.5)
        canvas.drawString(15 * mm, 10 * mm, f"{doc.get('doc_code', doc['id'])} · {doc['title'][:70]} · {BANNER_SHORT}")
        canvas.drawRightString(w - 15 * mm, 10 * mm, f"Page {d.page}")
        canvas.restoreState()

    story = [Paragraph(doc["title"], h1), Paragraph(_meta_line(doc), meta)]
    for blk in doc["blocks"]:
        kind = blk[0]
        if kind == "h":
            story.append(Paragraph(blk[1], h2))
        elif kind == "p":
            story.append(Paragraph(blk[1], body))
        elif kind == "note":
            story.append(Paragraph(blk[1], note))
            story.append(Spacer(1, 4))
        elif kind == "b":
            story.append(ListFlowable([ListItem(Paragraph(x, body), leftIndent=10) for x in blk[1]],
                                      bulletType="bullet", start="•", leftIndent=12))
        elif kind == "t":
            headers, rows = blk[1], blk[2]
            data = [[Paragraph(f"<b>{c}</b>", cell) for c in headers]] + [[Paragraph(str(c), cell) for c in r] for r in rows]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(KeepTogether(t))
            story.append(Spacer(1, 6))
    pdf = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=22 * mm, bottomMargin=16 * mm, title=doc["title"],
                            author="Fictional — educational prototype", subject=BANNER_SHORT)
    pdf.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


# ───────────────────────────── DOCX ─────────────────────────────
def render_docx(doc: dict, out: Path) -> None:
    d = Document()
    for section in d.sections:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(BANNER)
        run.bold = True
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
        fp = section.footer.paragraphs[0]
        fp.text = f"{doc.get('doc_code', doc['id'])} · {BANNER_SHORT}"
        fp.runs[0].font.size = Pt(7.5)
    d.add_heading(doc["title"], level=0)
    m = d.add_paragraph(_meta_line(doc))
    m.runs[0].font.size = Pt(8)
    for blk in doc["blocks"]:
        kind = blk[0]
        if kind == "h":
            d.add_heading(blk[1], level=1)
        elif kind == "p":
            d.add_paragraph(_strip_tags(blk[1]))
        elif kind == "note":
            p = d.add_paragraph()
            r = p.add_run(_strip_tags(blk[1]))
            r.italic = True
        elif kind == "b":
            for x in blk[1]:
                d.add_paragraph(_strip_tags(x), style="List Bullet")
        elif kind == "t":
            headers, rows = blk[1], blk[2]
            t = d.add_table(rows=1, cols=len(headers))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(headers):
                t.rows[0].cells[i].text = str(h)
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = _strip_tags(str(c))
            d.add_paragraph()
    d.core_properties.title = doc["title"]
    d.core_properties.subject = BANNER_SHORT
    d.core_properties.author = "Fictional — educational prototype"
    d.save(str(out))


def _strip_tags(s: str) -> str:
    import re
    return re.sub(r"</?(b|i|u|br/?)>", "", s)


# ───────────────────────────── XLSX ─────────────────────────────
def render_xlsx(doc: dict, out: Path) -> None:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = BANNER_SHORT
    readme["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    readme["A1"].fill = PatternFill("solid", fgColor="B00020")
    readme["A3"] = doc["title"]
    readme["A3"].font = Font(bold=True, size=13)
    readme["A4"] = _meta_line(doc)
    readme["A6"] = BANNER
    readme["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    readme.row_dimensions[6].height = 45
    for i, line in enumerate(doc.get("description", []), start=8):
        readme.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 120

    for sheet in doc["sheets"]:
        name, headers, rows = sheet[0], sheet[1], sheet[2]
        widths = sheet[3] if len(sheet) > 3 else None
        ws = wb.create_sheet(name[:31])
        ws["A1"] = BANNER_SHORT + f" · {doc['title']}"
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="B00020")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 4))
        for j, h in enumerate(headers, start=1):
            c = ws.cell(row=3, column=j, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8E8E8")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, r in enumerate(rows, start=4):
            for j, v in enumerate(r, start=1):
                ws.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths else 22)
        ws.freeze_panes = "A4"
    wb.properties.title = doc["title"]
    wb.properties.subject = BANNER_SHORT
    wb.save(str(out))


# ───────────────────────────── MD ─────────────────────────────
def render_md(doc: dict, out: Path) -> None:
    lines = [f"> **{BANNER}**", "", f"# {doc['title']}", "", f"_{_meta_line(doc)}_", ""]
    for blk in doc["blocks"]:
        kind = blk[0]
        if kind == "h":
            lines += [f"## {blk[1]}", ""]
        elif kind == "p":
            lines += [_md_inline(blk[1]), ""]
        elif kind == "note":
            lines += [f"> {_md_inline(blk[1])}", ""]
        elif kind == "b":
            lines += [f"- {_md_inline(x)}" for x in blk[1]] + [""]
        elif kind == "t":
            headers, rows = blk[1], blk[2]
            lines += ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
            lines += ["| " + " | ".join(_md_inline(str(c)) for c in r) + " |" for r in rows] + [""]
    out.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def _md_inline(s: str) -> str:
    return s.replace("<b>", "**").replace("</b>", "**").replace("<i>", "_").replace("</i>", "_").replace("<br/>", " ")


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "xlsx": render_xlsx, "md": render_md}
